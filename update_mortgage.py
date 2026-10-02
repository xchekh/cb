"""Обновление файлов статистики ЦБ РФ по ипотеке (региональный разрез).

Источник: https://www.cbr.ru/statistics/bank_sector/mortgage/#a_62951

Скрипт берёт свежие файлы ЦБ (скачивает сам или берёт уже скачанные)
и приводит их к виду, принятому в репозитории:
  * листы «в рублях», «в инвалюте», «итого» идут в этом порядке
    (у ЦБ «итого» первый);
  * на листах «итого», «срок в руб», «ставка в рублях» после столбца A
    вставлен столбец B со стандартными названиями регионов, а в строке
    заголовка — сквозная нумерация столбцов (1, =B2+1, =C2+1, ...).

Данные берутся из файла ЦБ целиком, потому что ЦБ пересматривает
прошлые месяцы.

Примеры:
  python update_mortgage.py                  # скачать с cbr.ru и перезаписать файлы в репо
  python update_mortgage.py --src ~/Downloads  # взять вручную скачанные файлы
  python update_mortgage.py --out new/       # записать результат в другую папку
  python update_mortgage.py --commit         # то же + git commit, если данные изменились (для cron)

Если данные у ЦБ не изменились, файлы не перезаписываются.
"""
import argparse
import copy
from datetime import datetime
import glob
import os
import subprocess
import sys
import tempfile
import urllib.request

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

BASE_URL = "https://www.cbr.ru/vfs/statistics/BankSector/Mortgage/"
REPO = os.path.dirname(os.path.abspath(__file__))

# файл ЦБ -> путь в репозитории
FILES = {
    "02_10_Quantity_mortgage.xlsx": "02_10_Quantity_mortgage (6).xlsx",
    "02_11_New_loans_mortgage.xlsx": "02_11_New_loans_mortgage (6).xlsx",
    "02_13_Rates_mortgage.xlsx": "02_13_Rates_mortgage (6).xlsx",
    "02_15_Quantity_scpa_mortgage.xlsx": "ДДУ/02_15_Quantity_scpa_mortgage (6).xlsx",
    "02_16_New_loans_scpa_mortgage.xlsx": "ДДУ/02_16_New_loans_scpa_mortgage (6) (1).xlsx",
    "02_17_Rates_scpa_mortgage.xlsx": "ДДУ/02_17_Rates_scpa_mortgage (6).xlsx",
}

# листы, в которые вставляется столбец B со справочником регионов
SHEETS_WITH_REGION_COL = {"итого", "срок в руб", "ставка в рублях"}
SHEET_ORDER = ["в рублях", "в инвалюте", "итого"]

# название в файле ЦБ -> стандартное название (столбец B)
REGIONS = {
    'РОССИЙСКАЯ ФЕДЕРАЦИЯ': 'Всего РФ',
    'ЦЕНТРАЛЬНЫЙ ФЕДЕРАЛЬНЫЙ ОКРУГ': None,
    'Белгородская область': 'Белгородская область',
    'Брянская область': 'Брянская область',
    'Владимирская область': 'Владимирская область',
    'Воронежская область': 'Воронежская область',
    'Ивановская область': 'Ивановская область',
    'Калужская область': 'Калужская область',
    'Костромская область': 'Костромская область',
    'Курская область': 'Курская область',
    'Липецкая область': 'Липецкая область',
    'Московская область': 'Московская область',
    'Орловская область': 'Орловская область',
    'Рязанская область': 'Рязанская область',
    'Смоленская область': 'Смоленская область',
    'Тамбовская область': 'Тамбовская область',
    'Тверская область': 'Тверская область',
    'Тульская область': 'Тульская область',
    'Ярославская область': 'Ярославская область',
    'г. Москва': 'Город Москва',
    'СЕВЕРО-ЗАПАДНЫЙ ФЕДЕРАЛЬНЫЙ ОКРУГ': None,
    'Республика Карелия': 'Республика Карелия',
    'Республика Коми': 'Республика Коми',
    'Архангельская область': 'Архангельская область',
    'в том числе Ненецкий автономный округ': None,
    'Архангельская область без данных по Ненецкому автономному округу': None,
    'Вологодская область': 'Вологодская область',
    'Калининградская область': 'Калининградская область',
    'Ленинградская область': 'Ленинградская область',
    'Мурманская область': 'Мурманская область',
    'Новгородская область': 'Новгородская область',
    'Псковская область': 'Псковская область',
    'г. Санкт-Петербург': 'Город Санкт-Петербург',
    'ЮЖНЫЙ ФЕДЕРАЛЬНЫЙ ОКРУГ': None,
    'Республика Адыгея (Адыгея)': 'Республика Адыгея',
    'Республика Калмыкия': 'Республика Калмыкия',
    'Республика Крым': 'Республика Крым',
    'Краснодарский край': 'Краснодарский край',
    'Астраханская область': 'Астраханская область',
    'Волгоградская область': 'Волгоградская область',
    'Ростовская область': 'Ростовская область',
    'г. Севастополь': 'Город Севастополь',
    'СЕВЕРО-КАВКАЗСКИЙ ФЕДЕРАЛЬНЫЙ ОКРУГ': None,
    'Республика Дагестан': 'Республика Дагестан',
    'Республика Ингушетия': 'Республика Ингушетия',
    'Кабардино-Балкарская Республика': 'Кабардино-Балкарская Республика',
    'Карачаево-Черкесская Республика': 'Карачаево-Черкесская Республика',
    'Республика Северная Осетия - Алания': 'Республика Северная Осетия',
    'Чеченская Республика': 'Чеченская Республика',
    'Ставропольский край': 'Ставропольский край',
    'ПРИВОЛЖСКИЙ ФЕДЕРАЛЬНЫЙ ОКРУГ': None,
    'Республика Башкортостан': 'Республика Башкортостан',
    'Республика Марий Эл': 'Республика Марий Эл',
    'Республика Мордовия': 'Республика Мордовия',
    'Республика Татарстан (Татарстан)': 'Республика Татарстан',
    'Удмуртская Республика': 'Удмуртская Республика',
    'Чувашская Республика - Чувашия': 'Чувашская Республика',
    'Пермский край': 'Пермский край',
    'Кировская область': 'Кировская область',
    'Нижегородская область': 'Нижегородская область',
    'Оренбургская область': 'Оренбургская область',
    'Пензенская область': 'Пензенская область',
    'Самарская область': 'Самарская область',
    'Саратовская область': 'Саратовская область',
    'Ульяновская область': 'Ульяновская область',
    'УРАЛЬСКИЙ ФЕДЕРАЛЬНЫЙ ОКРУГ': '#N/A',
    'Курганская область': 'Курганская область',
    'Свердловская область': 'Свердловская область',
    'Тюменская область': None,
    'в том числе Ханты-Мансийский автономный округ - Югра': 'Ханты-Мансийский АО - Югра',
    'в том числе Ямало-Ненецкий автономный округ': 'Ямало-Ненецкий АО',
    'Тюменская область без данных по Ханты-Мансийскому автономному округу - Югре и Ямало-Ненецкому автономному округу': 'Тюменская область',
    'Челябинская область': 'Челябинская область',
    'СИБИРСКИЙ ФЕДЕРАЛЬНЫЙ ОКРУГ': '#N/A',
    'Республика Алтай': 'Республика Алтай',
    'Республика Тыва': 'Республика Тыва',
    'Республика Хакасия': 'Республика Хакасия',
    'Алтайский край': 'Алтайский край',
    'Красноярский край': 'Красноярский край',
    'Иркутская область': 'Иркутская область',
    'Кемеровская область - Кузбасс': 'Кемеровская область - Кузбасс',
    'Новосибирская область': 'Новосибирская область',
    'Омская область': 'Омская область',
    'Томская область': 'Томская область',
    'ДАЛЬНЕВОСТОЧНЫЙ ФЕДЕРАЛЬНЫЙ ОКРУГ': None,
    'Республика Бурятия': 'Республика Бурятия',
    'Республика Саха (Якутия)': 'Республика Саха (Якутия)',
    'Забайкальский край': 'Забайкальский край',
    'Камчатский край': 'Камчатский край',
    'Приморский край': 'Приморский край',
    'Хабаровский край': 'Хабаровский край',
    'Амурская область': 'Амурская область',
    'Магаданская область': 'Магаданская область',
    'Сахалинская область': 'Сахалинская область',
    'Еврейская автономная область': 'Еврейская АО',
    'Чукотский автономный округ': None,
}


def download(dest):
    for name in FILES:
        req = urllib.request.Request(BASE_URL + name, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=60) as r, open(os.path.join(dest, name), "wb") as f:
            f.write(r.read())
        print(f"скачан {name}")


def find_source(src_dir, name):
    """Ищет файл по префиксу (02_10_...), т.к. браузер добавляет « (1)» и т.п."""
    exact = os.path.join(src_dir, name)
    if os.path.exists(exact):
        return exact
    prefix = name[:5]
    stem = os.path.splitext(name)[0]
    found = sorted(glob.glob(os.path.join(src_dir, stem + "*.xlsx")) or
                   glob.glob(os.path.join(src_dir, prefix + "*.xlsx")), key=os.path.getmtime)
    if not found:
        sys.exit(f"не найден файл {name} в {src_dir}")
    return found[-1]  # самый свежий


def title_row(ws):
    """Строка с названием таблицы: 2 у основных файлов, 1 у ДДУ."""
    for r in (1, 2, 3):
        if isinstance(ws.cell(r, 1).value, str):
            return r
    raise ValueError(f"не найден заголовок на листе {ws.title}")


def add_region_column(ws):
    t = title_row(ws)
    max_col = ws.max_column  # последний месяц у ЦБ

    for rng in list(ws.merged_cells.ranges):
        if rng.min_row == t:
            ws.unmerge_cells(str(rng))
    al = ws.cell(t, 1).alignment
    ws.cell(t, 1).alignment = Alignment(vertical=al.vertical, wrap_text=al.wrap_text)

    # ширины столбцов сдвигаем вручную: insert_cols их не трогает
    dims = {k: (d.min, d.max, d.width) for k, d in ws.column_dimensions.items()}
    ws.insert_cols(2)
    for k in list(ws.column_dimensions):
        if k != "A":
            del ws.column_dimensions[k]
    for k, (mn, mx, w) in dims.items():
        if k == "A":
            continue
        letter = get_column_letter(mn + 1)
        ws.column_dimensions[letter].width = w
        ws.column_dimensions[letter].min = mn + 1
        ws.column_dimensions[letter].max = mx + 1

    # столбец B: стиль как у A, значения из справочника
    unknown = []
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1), ws.cell(r, 2)
        b._style = copy.copy(a._style)
        if r > t + 1 and a.value is not None:
            if a.value not in REGIONS:
                unknown.append(a.value)
            b.value = REGIONS.get(a.value)

    # строка заголовка: 1, =B{t}+1, =C{t}+1, ... — стиль как у названия
    title_style = ws.cell(t, 1)
    for c in range(2, max_col + 2):
        cell = ws.cell(t, c)
        cell.value = "1" if c == 2 else f"={get_column_letter(c - 1)}{t}+1"
        cell.font = copy.copy(title_style.font)
        cell.fill = copy.copy(title_style.fill)
        cell.alignment = copy.copy(title_style.alignment)
        cell.number_format = "0.00"

    if ws.freeze_panes:
        fr = ws[ws.freeze_panes]
        ws.freeze_panes = ws.cell(fr.row, fr.column + 1).coordinate
    return unknown


def convert(src):
    """Читает файл ЦБ и приводит к виду репозитория. Возвращает (книга, последний месяц, новые регионы)."""
    wb = openpyxl.load_workbook(src)
    unknown = []
    for ws in wb.worksheets:
        if ws.title in SHEETS_WITH_REGION_COL:
            unknown += add_region_column(ws)
    if set(wb.sheetnames) == set(SHEET_ORDER):
        wb._sheets = [wb[n] for n in SHEET_ORDER]
        wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws is wb.active

    ws = wb.worksheets[0]
    last = ws.cell(title_row(ws) + 1, ws.max_column).value
    return wb, last, sorted(set(unknown))


def values(wb):
    return {ws.title: [tuple(c.value for c in row) for row in ws.iter_rows()] for ws in wb.worksheets}


def same_as_existing(wb, path):
    if not os.path.exists(path):
        return False
    return values(wb) == values(openpyxl.load_workbook(path))


def git(*args):
    return subprocess.run(["git", "-C", REPO, *args], check=True, capture_output=True, text=True).stdout


def commit(paths, last):
    rel = [os.path.relpath(p, REPO) for p in paths]
    git("add", "--", *rel)
    git("commit", "-m", f"Данные ЦБ по ипотеке: обновление по {last}", "--", *rel)
    print(f"закоммичено: {git('log', '-1', '--format=%h %s').strip()}")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--src", help="папка с уже скачанными файлами ЦБ (по умолчанию скачать с cbr.ru)")
    p.add_argument("--out", default=REPO, help="куда сохранить (по умолчанию — поверх файлов в репозитории)")
    p.add_argument("--commit", action="store_true", help="закоммитить изменённые файлы (только при --out по умолчанию)")
    args = p.parse_args()
    if args.commit and os.path.abspath(args.out) != REPO:
        sys.exit("--commit работает только при записи в репозиторий")

    print(f"=== {datetime.now():%Y-%m-%d %H:%M}")
    # сначала конвертируем всё в памяти: если ЦБ отдал битый файл, ничего не перезапишем
    with tempfile.TemporaryDirectory() as tmp:
        src_dir = args.src or tmp
        if not args.src:
            download(tmp)
        books = {target: convert(find_source(src_dir, name)) for name, target in FILES.items()}

    changed = []
    for target, (wb, last, unknown) in books.items():
        dst = os.path.join(args.out, target)
        for u in unknown:
            print(f"  ВНИМАНИЕ: нового региона «{u}» нет в справочнике REGIONS, столбец B пустой")
        if same_as_existing(wb, dst):
            print(f"{os.path.basename(dst)}: без изменений (по {last})")
            continue
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        wb.save(dst)
        changed.append(dst)
        print(f"{os.path.basename(dst)}: обновлён, данные по {last}")

    if changed and args.commit:
        commit(changed, last)


if __name__ == "__main__":
    main()
